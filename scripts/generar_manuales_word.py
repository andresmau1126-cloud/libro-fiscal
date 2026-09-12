from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    for name, size, color in (("Title", 22, "1F4E79"), ("Heading 1", 15, "1F4E79"), ("Heading 2", 12, "3F6B8A")):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Consolas"
        code_style.font.size = Pt(8.5)
        code_style.paragraph_format.left_indent = Inches(0.25)
        code_style.paragraph_format.space_after = Pt(6)


def add_inline(paragraph, text):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def add_table(document, rows):
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, rows[0]):
        cell.text = value.strip()
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows[1:]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value.strip()
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        from docx.oxml import OxmlElement
        shading = OxmlElement("w:shd")
        shading.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill", "1F4E79")
        tc_pr.append(shading)
    document.add_paragraph()


def convert(markdown_path, output_path):
    document = Document()
    configure_document(document)
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    code_lines = []

    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph(style="Code Block")
                paragraph.add_run("\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            paragraph = document.add_paragraph(style="Title" if level == 1 else f"Heading {level}")
            add_inline(paragraph, heading.group(2))
            index += 1
            continue

        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                values = [value.strip() for value in lines[index].strip().strip("|").split("|")]
                if not all(set(value) <= {"-", ":", " "} for value in values):
                    rows.append(values)
                index += 1
            if rows:
                add_table(document, rows)
            continue

        list_match = re.match(r"^[-*]\s+(.+)$", line)
        numbered_match = re.match(r"^\d+\.\s+(.+)$", line)
        if list_match or numbered_match:
            text = (list_match or numbered_match).group(1)
            paragraph = document.add_paragraph(style="List Bullet" if list_match else "List Number")
            add_inline(paragraph, text)
            index += 1
            continue

        paragraph = document.add_paragraph()
        add_inline(paragraph, line)
        index += 1

    document.core_properties.title = markdown_path.stem.replace("_", " ").title()
    document.core_properties.subject = "Libro Fiscal v2"
    document.save(output_path)


if __name__ == "__main__":
    convert(DOCS / "manual_tecnico.md", DOCS / "manual_tecnico.docx")
    convert(DOCS / "manual_usuario.md", DOCS / "manual_usuario.docx")
    print("Manuales Word generados en docs/")
